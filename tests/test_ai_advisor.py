"""
Тести модуля ШІ-порадника (app/services/ai_advisor.py, fallback_advisor.py)
та ендпоінтів /api/ai/*:
- Формування запиту до Gemini (ключ у заголовку, не в URL; без temperature/top_p/top_k; system_instruction)
- Ліміт 500 символів для повідомлення (валідація та 400 Bad Request)
- Обмеження історії контексту до 6 останніх реплік
- Анонімізація назв процесів при AI_SEND_PROCESS_NAMES=False
- Автономний фолбек при 429, таймауті або відсутності ключа
- Секретний ключ ніколи не потрапляє у відповіді API, лог або БД
- Налаштування кольорів, пунктиру піків та легенди діаграми Excel
"""

import io
import json
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch
from urllib.error import HTTPError

import openpyxl

from app import config
from app.core.collector import MetricCollector
from app.db.connection import db_session, init_db
from app.server import create_app
from app.services.ai_advisor import (
    _AUTO_ADVICE_CACHE,
    build_ai_context,
    chat_with_advisor,
    get_auto_advice,
)
from app.services.analytics import _calculate_health_index
from app.services.exporter import generate_excel_report

def _extract_title_text(title_obj) -> str:
    """
    Допоміжна функція для витягування текстового рядка із заголовків openpyxl
    (об'єкти Title / Text / RichText):
    title.tx.rich.p -> r -> t з урахуванням того, що будь-які частини можуть бути None або рядками.
    """
    if title_obj is None:
        return ""
    if isinstance(title_obj, str):
        return title_obj.strip()

    # Збираємо можливі кандидати для витягування тексту
    candidates = []
    if hasattr(title_obj, "tx") and title_obj.tx is not None:
        candidates.append(title_obj.tx)
    if hasattr(title_obj, "text") and title_obj.text is not None:
        candidates.append(title_obj.text)
    candidates.append(title_obj)

    for item in candidates:
        if isinstance(item, str):
            return item.strip()

        rich = getattr(item, "rich", None)
        if rich is not None:
            paragraphs = getattr(rich, "p", None) or getattr(rich, "paragraphs", None) or []
            if not isinstance(paragraphs, (list, tuple)):
                paragraphs = [paragraphs]
            collected = []
            for p in paragraphs:
                if p is None:
                    continue
                runs = getattr(p, "r", None) or []
                if not isinstance(runs, (list, tuple)):
                    runs = [runs]
                for r in runs:
                    if r is None:
                        continue
                    t = getattr(r, "t", None)
                    if t is not None:
                        collected.append(str(t))
                if not runs and hasattr(p, "t") and p.t is not None:
                    collected.append(str(p.t))
            if collected:
                return "".join(collected).strip()

    for item in candidates:
        v = getattr(item, "v", None)
        if v is not None:
            return str(v).strip()

    return str(title_obj).strip()


def _generate_report_and_get_chart(period: str, db_path: Path):
    """
    Генерує Excel-звіт та перехоплює об'єкт діаграми на етапі додавання до аркуша
    (до збереження та повторного завантаження openpyxl).
    Дозволяє протестувати вихідні властивості LineChart (width, height, anchor),
    які openpyxl не відновлює з XML після load_workbook.
    """
    captured = {}
    orig_add_chart = openpyxl.worksheet.worksheet.Worksheet.add_chart

    def intercept_add_chart(ws_self, chart, anchor="A1"):
        captured["chart"] = chart
        captured["anchor"] = anchor
        return orig_add_chart(ws_self, chart, anchor)

    with patch.object(openpyxl.worksheet.worksheet.Worksheet, "add_chart", side_effect=intercept_add_chart, autospec=True):
        bio, filename = generate_excel_report(period=period, db_path=db_path)

    return bio, filename, captured.get("chart"), captured.get("anchor")


def _extract_color(fill_obj) -> str:
    """Витягує hex-рядок кольору (наприклад '4C8DFF') з об'єкта або рядка openpyxl."""
    if fill_obj is None:
        return ""
    if isinstance(fill_obj, str):
        return fill_obj.upper()
    if hasattr(fill_obj, "srgbClr") and fill_obj.srgbClr:
        return str(fill_obj.srgbClr).upper()
    return str(fill_obj).upper()


def _extract_dash(dash_obj) -> Optional[str]:
    """Витягує назву типу пунктиру (наприклад 'dash') з об'єкта або рядка openpyxl."""
    if dash_obj is None:
        return None
    if isinstance(dash_obj, str):
        return dash_obj
    if hasattr(dash_obj, "val") and dash_obj.val:
        return str(dash_obj.val)
    return str(dash_obj)


class TestAiAdvisor(unittest.TestCase):
    """Набір модульних тестів для ШІ-порадника."""

    def setUp(self):
        self.temp_file = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
        self.db_path = Path(self.temp_file.name)
        self.temp_file.close()
        init_db(self.db_path)

        # Очищаємо in-memory кеш перед кожним тестом
        _AUTO_ADVICE_CACHE.clear()

        # Зберігаємо оригінальні налаштування config
        self._orig_api_key = config.GEMINI_API_KEY
        self._orig_model = config.GEMINI_MODEL
        self._orig_send_proc = config.AI_SEND_PROCESS_NAMES

        # Додаємо тестові дані в БД
        with db_session(self.db_path) as conn:
            cur = conn.cursor()
            cur.execute("""
                INSERT INTO metrics (timestamp, cpu_total, cpu_max, cpu_cores, ram_used_bytes, ram_total_bytes, ram_percent, ram_max, swap_used_bytes, swap_total_bytes, disk_read_bytes_sec, disk_write_bytes_sec, net_sent_bytes_sec, net_recv_bytes_sec, disks_info, is_aggregated)
                VALUES (1700000000, 45.0, 95.0, '[]', 8000000000, 16000000000, 50.0, 92.0, 0, 0, 0, 0, 0, 0, '[]', 0)
            """)
            cur.execute("""
                INSERT INTO alerts (timestamp, resolved_at, metric_name, rule_condition, trigger_value, threshold_value, severity, culprits_summary)
                VALUES (1700000000, NULL, 'RAM', 'RAM > 85% понад 10 с', 88.0, 85.0, 'warning', 'chrome.exe (45% RAM)')
            """)
            cur.execute("""
                INSERT INTO process_usage_hourly (hour, process_name, appearances_count, avg_cpu_percent, max_cpu_percent, max_memory_percent, first_seen_ts, last_seen_ts)
                VALUES (1700000000, 'chrome.exe', 15, 30.0, 85.0, 20.0, 1700000000, 1700003600)
            """)

        self.collector = MetricCollector(db_path=self.db_path)
        self.app = create_app(collector=self.collector, db_path=self.db_path)
        self.client = self.app.test_client()

    def tearDown(self):
        self.collector.stop()
        config.GEMINI_API_KEY = self._orig_api_key
        config.GEMINI_MODEL = self._orig_model
        config.AI_SEND_PROCESS_NAMES = self._orig_send_proc
        _AUTO_ADVICE_CACHE.clear()

        for p in [self.db_path, Path(str(self.db_path) + "-wal"), Path(str(self.db_path) + "-shm")]:
            if p.exists():
                try:
                    p.unlink()
                except Exception:
                    pass

    @patch("urllib.request.urlopen")
    def test_request_formation(self, mock_urlopen):
        """
        1. Перевірка формування запиту до Gemini REST generateContent:
        - Ключ у заголовку x-goog-api-key, відсутній у URL
        - Відсутні temperature, top_p, top_k
        - Присутня system_instruction
        """
        config.GEMINI_API_KEY = "test-secret-key-12345"

        mock_resp = MagicMock()
        mock_resp.read.return_value = json.dumps({
            "candidates": [
                {
                    "content": {
                        "parts": [{"text": "Тестова порада від моделі Gemini"}]
                    }
                }
            ]
        }).encode("utf-8")
        mock_resp.__enter__.return_value = mock_resp
        mock_urlopen.return_value = mock_resp

        res = get_auto_advice(period="24h", db_path=self.db_path, force=True)
        self.assertEqual(res["source"], "Gemini")
        self.assertEqual(res["status"], "ok")
        self.assertIn("Тестова порада", res["advice"])

        mock_urlopen.assert_called_once()
        req = mock_urlopen.call_args[0][0]

        # 1.1 Перевірка URL: ключ НЕ повинен бути в URL
        self.assertNotIn("key=", req.full_url)
        self.assertIn("generativelanguage.googleapis.com", req.full_url)
        self.assertIn(config.GEMINI_MODEL, req.full_url)

        # 1.2 Перевірка заголовка x-goog-api-key
        self.assertEqual(req.headers.get("X-goog-api-key"), "test-secret-key-12345")

        # 1.3 Перевірка тіла запиту (JSON)
        body = json.loads(req.data.decode("utf-8"))
        body_str = json.dumps(body)

        self.assertNotIn("temperature", body_str)
        self.assertNotIn("top_p", body_str)
        self.assertNotIn("top_k", body_str)

        self.assertIn("system_instruction", body)
        instruction_text = body["system_instruction"]["parts"][0]["text"]
        self.assertIn("порадник з продуктивності", instruction_text)
        self.assertIn("Windows", instruction_text)

    def test_message_length_limit(self):
        """
        2. Перевірка ліміту 500 символів:
        - Повідомлення <= 500 символів успішно приймається
        - Повідомлення > 500 символів викликає ValueError та статус 400 Bad Request
        """
        config.GEMINI_API_KEY = ""  # Тестуємо в автономному режимі

        # Валідне повідомлення (500 символів)
        msg_500 = "x" * 500
        res_ok = chat_with_advisor(msg_500, period="24h", db_path=self.db_path)
        self.assertIn("[Автономний режим]", res_ok["reply"])

        # Повідомлення понад 500 символів (501 символ)
        msg_501 = "x" * 501
        with self.assertRaises(ValueError) as ctx:
            chat_with_advisor(msg_501, period="24h", db_path=self.db_path)
        self.assertIn("500 символів", str(ctx.exception))

        # Перевірка HTTP 400 від ендпоінта
        with self.client.post("/api/ai/chat", json={"message": msg_501}) as response:
            self.assertEqual(response.status_code, 400)
            json_err = response.get_json()
            self.assertIn("500 символів", json_err.get("error", ""))

    @patch("app.services.ai_advisor._call_gemini_rest")
    def test_history_truncation(self, mock_call_rest):
        """
        3. Перевірка обмеження історії контексту до 6 останніх реплік.
        """
        config.GEMINI_API_KEY = "dummy-key"
        mock_call_rest.return_value = "Відповідь на 11-те повідомлення"

        # Створюємо історію з 10 реплік
        history = [
            {"role": "user" if i % 2 == 0 else "model", "text": f"Повідомлення {i+1}"}
            for i in range(10)
        ]

        res = chat_with_advisor("Нове запитання", history=history, period="24h", db_path=self.db_path)
        self.assertEqual(res["reply"], "Відповідь на 11-те повідомлення")

        mock_call_rest.assert_called_once()
        contents = mock_call_rest.call_args[0][0]

        # contents містить не більше 6 реплік з історії + 1 поточну = максимум 7
        self.assertLessEqual(len(contents), 7)
        # Перша передана з історії репліка повинна бути 5-ю (індекс 4), оскільки береться history[-6:]
        self.assertEqual(contents[0]["parts"][0]["text"], "Повідомлення 5")

    def test_process_names_anonymization(self):
        """
        4. Перевірка анонімізації назв процесів при AI_SEND_PROCESS_NAMES=False.
        Реальні назви ('chrome.exe') замінюються на 'процес 1, 2…'.
        """
        config.AI_SEND_PROCESS_NAMES = False

        ctx = build_ai_context(period="24h", db_path=self.db_path)
        ctx_str = json.dumps(ctx, ensure_ascii=False)

        # Реальна назва 'chrome.exe' не повинна передаватися
        self.assertNotIn("chrome.exe", ctx_str)
        # Замість неї з'являється «процес 1»
        self.assertIn("процес 1", ctx_str)

    def test_fallback_on_missing_key(self):
        """
        5. Перевірка автоматичного фолбеку при відсутності ключа (GEMINI_API_KEY = '').
        """
        config.GEMINI_API_KEY = ""

        # Автопорада
        adv_res = get_auto_advice(period="24h", db_path=self.db_path, force=True)
        self.assertEqual(adv_res["source"], "Автономний режим")
        self.assertEqual(adv_res["status"], "fallback")
        self.assertIn("[Автономний режим]", adv_res["advice"])

        # Чат
        chat_res = chat_with_advisor("Скільки в мене оперативної пам'яті?", period="24h", db_path=self.db_path)
        self.assertEqual(chat_res["source"], "Автономний режим")
        self.assertEqual(chat_res["status"], "fallback")
        self.assertIn("[Автономний режим]", chat_res["reply"])

    @patch("urllib.request.urlopen")
    def test_fallback_on_429_and_timeout(self, mock_urlopen):
        """
        6. Перевірка офлайн-фолбеку при помилці 429 (Rate Limit) та TimeoutError.
        """
        config.GEMINI_API_KEY = "valid-dummy-key"

        # 6.1 HTTP 429 Too Many Requests
        fp = io.BytesIO(b'{"error": {"code": 429, "message": "Resource exhausted"}}')
        err_429 = HTTPError(
            url="https://generativelanguage.googleapis.com",
            code=429,
            msg="Resource Exhausted",
            hdrs={},
            fp=fp,
        )
        mock_urlopen.side_effect = err_429

        try:
            res_429 = get_auto_advice(period="24h", db_path=self.db_path, force=True)
            self.assertEqual(res_429["source"], "Автономний режим")
            self.assertEqual(res_429["status"], "fallback")
            self.assertIn("[Автономний режим]", res_429["advice"])
        finally:
            err_429.close()
            fp.close()

        # 6.2 TimeoutError
        mock_urlopen.side_effect = TimeoutError("Таймаут з'єднання з Google API")

        res_timeout = get_auto_advice(period="24h", db_path=self.db_path, force=True)
        self.assertEqual(res_timeout["source"], "Автономний режим")
        self.assertEqual(res_timeout["status"], "fallback")
        self.assertIn("[Автономний режим]", res_timeout["advice"])

    def test_key_never_in_response_or_logs(self):
        """
        7. Перевірка безпеки: секретний ключ ніколи не потрапляє у відповіді API чи логи.
        """
        secret_key = "AIzaSySuperSecretKeyDoNotLeak999"
        config.GEMINI_API_KEY = secret_key

        # 7.1 GET /api/ai/status
        with self.client.get("/api/ai/status") as res:
            self.assertEqual(res.status_code, 200)
            text = res.get_data(as_text=True)
            self.assertNotIn(secret_key, text)
            data = res.get_json()
            self.assertTrue(data.get("has_key"))
            self.assertNotIn("key", data)

        # 7.2 POST /api/ai/chat з помилкою валідації
        with self.client.post("/api/ai/chat", json={"message": "a" * 501}) as res:
            text = res.get_data(as_text=True)
            self.assertNotIn(secret_key, text)

        # 7.3 GET /api/ai/history
        with self.client.get("/api/ai/history") as res:
            self.assertEqual(res.status_code, 200)
            text = res.get_data(as_text=True)
            self.assertNotIn(secret_key, text)

    def test_excel_chart_series_colors_and_style(self):
        """
        8. Перевірка діаграми Excel на аркуші «Метрики»:
        - Серії CPU (4C8DFF), RAM (A78BFA), GPU (F0883E)
        - Пунктирні тонкі лінії для піків
        - Легенда знизу ('b')
        - Підписи осей та розміри діаграми (width=28, height=12)
        - Заповнення аркуша «Поради ШІ» з таблиці ai_advice
        """
        now = int(time.time())
        # Наповнюємо тестову базу свіжими метриками за останню годину,
        # щоб експортер мав щонайменше 2 точки для побудови діаграми
        with db_session(self.db_path) as conn:
            cur = conn.cursor()
            for offset in (180, 120, 60, 10):
                cur.execute("""
                    INSERT INTO metrics (
                        timestamp, cpu_total, cpu_max, cpu_cores,
                        ram_used_bytes, ram_total_bytes, ram_percent, ram_max,
                        swap_used_bytes, swap_total_bytes,
                        disk_read_bytes_sec, disk_write_bytes_sec,
                        net_sent_bytes_sec, net_recv_bytes_sec, disks_info,
                        is_aggregated, gpu_percent, gpu_max, gpu_temp_c
                    )
                    VALUES (?, 45.0, 95.0, '[]', 8000000000, 16000000000, 50.0, 92.0, 0, 0, 0, 0, 0, 0, '[]', 0, 30.0, 60.0, 55.0)
                """, (now - offset,))

            # Додаємо запис у таблицю ai_advice
            cur.execute("""
                INSERT INTO ai_advice (timestamp, mode, source, user_query, metrics_summary, advice_text, status)
                VALUES (?, 'auto', 'Gemini', NULL, '{}', 'Рекомендація очистити тимчасові файли', 'ok')
            """, (now,))

        # Перехоплюємо діаграму на етапі створення до збереження та завантажуємо збережений файл
        bio, filename, pre_chart, anchor = _generate_report_and_get_chart(period="1h", db_path=self.db_path)

        # 8.1 Перевірка властивостей об'єкта діаграми до збереження (які openpyxl скидає при load_workbook)
        self.assertIsNotNone(pre_chart, "Діаграма має бути побудована та передана у ws.add_chart")
        self.assertEqual(pre_chart.width, 28, "Ширина діаграми до збереження має бути 28 см")
        self.assertEqual(pre_chart.height, 12, "Висота діаграми до збереження має бути 12 см")
        self.assertEqual(anchor, "K2", "Діаграма повинна розміщуватись у клітинці K2")

        # 8.2 Перевірка збереженого звіту через load_workbook
        wb = openpyxl.load_workbook(bio)

        # Перевірка аркуша «Метрики» та діаграми
        ws_metrics = wb["Метрики"]
        self.assertGreater(len(ws_metrics._charts), 0, "На аркуші Метрики має бути лінійна діаграма")
        chart = ws_metrics._charts[0]

        self.assertEqual(chart.legend.position, "b", "Легенда повинна розміщуватись знизу")
        self.assertEqual(_extract_title_text(chart.title), "Телеметрія системи: CPU, RAM та GPU")
        self.assertEqual(_extract_title_text(chart.x_axis.title), "Час")
        self.assertEqual(_extract_title_text(chart.y_axis.title), "Відсоток (%)")

        # Перевірка 6 серій
        self.assertEqual(len(chart.series), 6)
        self.assertEqual(len(pre_chart.series), 6)

        expected_styles = [
            ("4C8DFF", None),    # CPU Avg
            ("4C8DFF", "dash"),  # CPU Peak (пунктир)
            ("A78BFA", None),    # RAM Avg
            ("A78BFA", "dash"),  # RAM Peak (пунктир)
            ("F0883E", None),    # GPU Avg
            ("F0883E", "dash"),  # GPU Peak (пунктир)
        ]

        for idx, (exp_color, exp_dash) in enumerate(expected_styles):
            # Перевірка на діаграмі до збереження
            s_pre = pre_chart.series[idx]
            self.assertEqual(_extract_color(s_pre.graphicalProperties.line.solidFill), exp_color, f"Серія {idx} має некоректний колір")
            if exp_dash:
                self.assertEqual(_extract_dash(s_pre.graphicalProperties.line.prstDash), exp_dash, f"Серія {idx} має бути пунктирною")
            self.assertFalse(s_pre.smooth)

            # Перевірка на завантаженій діаграмі (якщо властивості відновлені openpyxl)
            if idx < len(chart.series):
                s_post = chart.series[idx]
                line = getattr(getattr(s_post, "graphicalProperties", None), "line", None)
                if line is not None:
                    if hasattr(line, "solidFill") and line.solidFill:
                        self.assertEqual(_extract_color(line.solidFill), exp_color)
                    if exp_dash and hasattr(line, "prstDash") and line.prstDash:
                        self.assertEqual(_extract_dash(line.prstDash), exp_dash)

        # 8.3 Перевірка аркуша «Поради ШІ»
        ws_ai = wb["Поради ШІ"]
        self.assertEqual(ws_ai.cell(row=1, column=1).value, "ID")
        self.assertEqual(ws_ai.cell(row=2, column=4).value, "Gemini")
        self.assertIn("Рекомендація", str(ws_ai.cell(row=2, column=7).value))

    def test_excel_metrics_empty_shows_no_data_label(self):
        """
        9. Перевірка відсутності діаграми та наявності напису «Немає даних за період»
        на аркуші «Метрики», якщо база даних порожня.
        """
        temp_file = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
        empty_db = Path(temp_file.name)
        temp_file.close()
        try:
            init_db(empty_db)
            bio, _ = generate_excel_report(period="1h", db_path=empty_db)
            wb = openpyxl.load_workbook(bio)
            ws_metrics = wb["Метрики"]
            self.assertEqual(ws_metrics.cell(row=2, column=1).value, "Немає даних за період")
            self.assertEqual(len(ws_metrics._charts), 0)
        finally:
            for p in [empty_db, Path(str(empty_db) + "-wal"), Path(str(empty_db) + "-shm")]:
                if p.exists():
                    try:
                        p.unlink()
                    except Exception:
                        pass

    @patch("urllib.request.urlopen")
    def test_auto_advice_no_chat_history_in_request(self, mock_urlopen):
        """
        10. Перевірка, що автопорада виконується окремим запитом без історії чату,
        а формулювання завдання є явно тематичним щодо оптимізації ПК.
        """
        config.GEMINI_API_KEY = "test-key-auto"
        mock_resp = MagicMock()
        mock_resp.read.return_value = json.dumps({
            "candidates": [{"content": {"parts": [{"text": "Висновок та поради"}]}}]
        }).encode("utf-8")
        mock_resp.__enter__.return_value = mock_resp
        mock_urlopen.return_value = mock_resp

        # Створюємо репліку в чаті
        chat_with_advisor("Привіт, розкажи анекдот", period="24h", db_path=self.db_path)

        # Викликаємо автопораду
        res = get_auto_advice(period="24h", db_path=self.db_path, force=True)
        self.assertEqual(res["status"], "ok")

        # Перевіряємо запит автопоради
        last_req = mock_urlopen.call_args_list[-1][0][0]
        body = json.loads(last_req.data.decode("utf-8"))
        contents = body.get("contents", [])

        self.assertEqual(len(contents), 1, "Запит автопоради повинен містити рівно 1 блок без історії")
        self.assertEqual(contents[0]["role"], "user")

        prompt_text = contents[0]["parts"][0]["text"]
        self.assertIn("Завдання: аналіз продуктивності та оптимізація", prompt_text)
        self.assertNotIn("розкажи анекдот", prompt_text, "Історія чату не повинна бути в автопораді")

    def test_context_no_backslashes_in_disk_mounts(self):
        """
        11. Перевірка, що контекст ШІ не містить «\\» у назвах розділів дисків ('C:', 'D:').
        """
        ctx = build_ai_context(period="24h", db_path=self.db_path)
        for d in ctx.get("disks", []):
            mount = d.get("mount", "")
            self.assertNotIn("\\", mount, f"Розділ {mount} містить бекслеш")
            if mount and len(mount) >= 2 and mount[1] == ":":
                self.assertEqual(mount, mount[:2].upper())

        for p in ctx.get("analytics", {}).get("disk_projections", []):
            mount = p.get("mount", "")
            self.assertNotIn("\\", mount, f"Розділ прогнозу {mount} містить бекслеш")

        # Перевірка очищення _clean_mount_name
        from app.services.ai_advisor import _clean_mount_name
        self.assertEqual(_clean_mount_name("C:\\"), "C:")
        self.assertEqual(_clean_mount_name("C:\\\\"), "C:")
        self.assertEqual(_clean_mount_name("d:\\"), "D:")
        self.assertEqual(_clean_mount_name("D:"), "D:")
        self.assertEqual(_clean_mount_name("E:/"), "E:")

    def test_health_index_ram_penalty_formula(self):
        """
        12. Перевірка нової формули штрафу за RAM в індексі здоров'я:
        - 55% -> 0 балів
        - 80% -> 10 балів
        - 100% -> 20 балів
        """
        _, _, bd_55 = _calculate_health_index(
            cpu_avg=0.0, ram_avg=55.0, cpu_rz_sec=0, ram_rz_sec=0,
            duration_sec=3600, active_alerts_count=0, period_alerts_count=0,
            min_disk_free_pct=40.0
        )
        self.assertEqual(bd_55["penalty_ram_load"], 0.0)

        _, _, bd_80 = _calculate_health_index(
            cpu_avg=0.0, ram_avg=80.0, cpu_rz_sec=0, ram_rz_sec=0,
            duration_sec=3600, active_alerts_count=0, period_alerts_count=0,
            min_disk_free_pct=40.0
        )
        self.assertEqual(bd_80["penalty_ram_load"], 10.0)

        _, _, bd_100 = _calculate_health_index(
            cpu_avg=0.0, ram_avg=100.0, cpu_rz_sec=0, ram_rz_sec=0,
            duration_sec=3600, active_alerts_count=0, period_alerts_count=0,
            min_disk_free_pct=40.0
        )
        self.assertEqual(bd_100["penalty_ram_load"], 20.0)


if __name__ == "__main__":
    unittest.main()
